"""Extract plain text from an uploaded CV (PDF, DOCX or TXT)."""
import io

MAX_PDF_PAGES = 10


def extract_cv_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES])

    elif name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)

    else:
        text = data.decode("utf-8", errors="ignore")

    return "\n".join(line.rstrip() for line in text.splitlines()).strip()